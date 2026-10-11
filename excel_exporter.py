"""
Excel export & bonus calculation for the Family Bonus Tracker.

Replaces the old Google Sheets integration. Bonuses are now calculated inside
the bot and stored in the local submissions store. An admin can request an
Excel export at any time via the !export command.
"""

import io
import logging
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config import BASE_ATTENDANCE_RATE, RATE_PER_KILL

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Bonus calculation
# ---------------------------------------------------------------------------
def calculate_bonus(event_type: str, kills: int) -> int:
    """
    Total Bonus = Base Attendance Bonus + (Kills * Rate Per Kill).

    Falls back to $0/kill if the event type is unknown.
    """
    per_kill = RATE_PER_KILL.get(event_type, 0)
    return BASE_ATTENDANCE_RATE + (kills * per_kill)


# ---------------------------------------------------------------------------
# Excel workbook builder
# ---------------------------------------------------------------------------
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
_RIGHT_ALIGN = Alignment(horizontal="right")


def _write_header(ws, headers: List[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws) -> None:
    """Best-effort column width based on cell content length."""
    for column_cells in ws.columns:
        max_len = 0
        col_letter = column_cells[0].column_letter
        for cell in column_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)


def build_export_workbook(submissions: List[dict]) -> io.BytesIO:
    """
    Build an in-memory .xlsx with two sheets:

      * **Summary** — one row per user: Username, User ID, Total Kills,
        Total Bonus, Events Attended. Sorted by Total Bonus descending.
      * **Details** — one row per approved participant: Date, Event Type,
        Submitter, User ID, Username, Kills, Bonus.

    Returns a BytesIO buffer ready to be attached to a Discord message.
    """
    approved = [s for s in submissions if s.get("status") == "approved"]

    wb = Workbook()

    # ---- Summary sheet ----
    ws_summary = wb.active
    ws_summary.title = "Summary"
    _write_header(
        ws_summary,
        ["Username", "User ID", "Total Kills", "Total Bonus", "Events Attended"],
    )

    totals: Dict[int, Dict[str, int]] = {}
    for sub in approved:
        event_type = sub.get("event_type", "")
        for p in sub.get("participants", []):
            uid = int(p["user_id"])
            kills = int(p.get("kills", 0))
            # Recompute bonus as a safety net in case old stored data lacks it.
            bonus = int(p.get("bonus", calculate_bonus(event_type, kills)))
            entry = totals.setdefault(
                uid,
                {
                    "username": p.get("username") or f"User#{uid}",
                    "kills": 0,
                    "bonus": 0,
                    "events": 0,
                },
            )
            entry["kills"] += kills
            entry["bonus"] += bonus
            entry["events"] += 1

    # Sort by total bonus descending.
    sorted_rows = sorted(totals.items(), key=lambda kv: kv[1]["bonus"], reverse=True)

    for idx, (uid, data) in enumerate(sorted_rows, start=2):
        ws_summary.cell(row=idx, column=1, value=data["username"])
        ws_summary.cell(row=idx, column=2, value=str(uid))
        ws_summary.cell(row=idx, column=3, value=data["kills"]).alignment = _RIGHT_ALIGN
        ws_summary.cell(row=idx, column=4, value=data["bonus"]).alignment = _RIGHT_ALIGN
        ws_summary.cell(row=idx, column=5, value=data["events"]).alignment = _RIGHT_ALIGN

    # Grand total row.
    if sorted_rows:
        total_row = len(sorted_rows) + 2
        ws_summary.cell(row=total_row, column=1, value="GRAND TOTAL").font = Font(bold=True)
        total_bonus = sum(d["bonus"] for _, d in sorted_rows)
        total_kills = sum(d["kills"] for _, d in sorted_rows)
        ws_summary.cell(row=total_row, column=3, value=total_kills).font = Font(bold=True)
        ws_summary.cell(row=total_row, column=4, value=total_bonus).font = Font(bold=True)

    _autosize(ws_summary)

    # ---- Details sheet ----
    ws_details = wb.create_sheet("Details")
    _write_header(
        ws_details,
        ["Date", "Event Type", "Submitter", "User ID", "Username", "Kills", "Bonus"],
    )

    row_idx = 2
    for sub in approved:
        event_type = sub.get("event_type", "")
        submitter = sub.get("submitter_name", "")
        date = sub.get("timestamp", "")
        for p in sub.get("participants", []):
            kills = int(p.get("kills", 0))
            bonus = int(p.get("bonus", calculate_bonus(event_type, kills)))
            ws_details.cell(row=row_idx, column=1, value=date)
            ws_details.cell(row=row_idx, column=2, value=event_type)
            ws_details.cell(row=row_idx, column=3, value=submitter)
            ws_details.cell(row=row_idx, column=4, value=str(p["user_id"]))
            ws_details.cell(row=row_idx, column=5, value=p.get("username") or f"User#{p['user_id']}")
            ws_details.cell(row=row_idx, column=6, value=kills).alignment = _RIGHT_ALIGN
            ws_details.cell(row=row_idx, column=7, value=bonus).alignment = _RIGHT_ALIGN
            row_idx += 1

    _autosize(ws_details)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    logger.info(
        "Built export workbook: %d approved submissions, %d unique users",
        len(approved),
        len(totals),
    )
    return buf
